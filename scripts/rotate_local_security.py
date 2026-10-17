"""Rotate local JWT and fictional demo credentials without printing secrets."""

from __future__ import annotations

import csv
import secrets
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from api.app.core.security import hash_password
from api.app.database import SessionLocal
from api.app.models import User
from api.app.services.demo_users import WORKBOOK_PATH


ROOT = Path(__file__).resolve().parents[1]
ENV_PATH = ROOT / ".env"
PRIVATE_DIRECTORY = ROOT / "data" / "private"


def replace_env_value(contents: str, key: str, value: str) -> str:
    lines = contents.splitlines()
    replacement = f"{key}={value}"
    for index, line in enumerate(lines):
        if line.startswith(f"{key}="):
            lines[index] = replacement
            break
    else:
        lines.append(replacement)
    return "\n".join(lines) + "\n"


def demo_usernames() -> list[str]:
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        sheet = workbook["Demo Users"]
        return [
            str(row[0]).strip().lower()
            for row in sheet.iter_rows(min_row=6, min_col=1, max_col=1, values_only=True)
            if row[0]
        ]
    finally:
        workbook.close()


def main() -> None:
    if not ENV_PATH.exists():
        raise SystemExit(".env was not found; refusing to create an incomplete production configuration.")

    replacements: list[tuple[str, str]] = []
    with SessionLocal() as db:
        for username in demo_usernames():
            user = db.scalar(select(User).where(User.username == username))
            if not user:
                continue
            password = secrets.token_urlsafe(24)
            user.password_hash = hash_password(password)
            replacements.append((username, password))
        db.commit()

    PRIVATE_DIRECTORY.mkdir(parents=True, exist_ok=True)
    credentials_path = PRIVATE_DIRECTORY / "rotated_demo_credentials.csv"
    with credentials_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["username", "password", "rotated_at"])
        rotated_at = datetime.now().astimezone().isoformat(timespec="seconds")
        writer.writerows((username, password, rotated_at) for username, password in replacements)

    env_contents = ENV_PATH.read_text(encoding="utf-8")
    env_contents = replace_env_value(env_contents, "JWT_SECRET_KEY", secrets.token_urlsafe(64))
    env_contents = replace_env_value(env_contents, "APP_DEBUG", "false")
    env_contents = replace_env_value(env_contents, "APP_ENV", "production")
    ENV_PATH.write_text(env_contents, encoding="utf-8")

    print(f"Rotated {len(replacements)} fictional demo accounts.")
    print("JWT secret rotated; debug disabled; production mode enabled.")
    print("Replacement credentials were written under data/private and were not displayed.")


if __name__ == "__main__":
    main()
