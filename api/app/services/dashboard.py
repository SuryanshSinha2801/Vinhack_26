from openpyxl import load_workbook

from api.app.schemas.dashboard import DailyWellbeingRecord, DashboardSummary, UserDashboard
from api.app.services.demo_users import WORKBOOK_PATH
from api.app.services.wellbeing_model import predict_wellbeing


def get_user_dashboard(username: str, display_name: str) -> UserDashboard | None:
    if not WORKBOOK_PATH.exists():
        return None
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    sheet = workbook["15-Day History"]
    history: list[DailyWellbeingRecord] = []
    for row in sheet.iter_rows(min_row=2, min_col=1, max_col=10, values_only=True):
        record_date, record_username, mood, stress, sleep, energy, connectedness, completed, support, note = row
        if str(record_username).strip().lower() != username.lower():
            continue
        history.append(
            DailyWellbeingRecord(
                date=record_date.date() if hasattr(record_date, "date") else record_date,
                mood=mood,
                stress=stress,
                sleep_hours=sleep,
                energy=energy,
                connectedness=connectedness,
                checkin_completed=bool(completed),
                support_response=str(support),
                note=str(note),
            )
        )
    workbook.close()
    if not history:
        return None
    history.sort(key=lambda item: item.date)
    latest = history[-1]
    count = len(history)
    summary = DashboardSummary(
        days=count,
        average_mood=round(sum(item.mood for item in history) / count, 1),
        average_stress=round(sum(item.stress for item in history) / count, 1),
        average_sleep_hours=round(sum(item.sleep_hours for item in history) / count, 1),
        latest_mood=latest.mood,
        latest_stress=latest.stress,
    )
    return UserDashboard(
        username=username,
        display_name=display_name,
        summary=summary,
        history=history,
        insight=predict_wellbeing(history),
    )
