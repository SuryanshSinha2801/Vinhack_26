import json
from datetime import date
from threading import Lock

from openpyxl import load_workbook

from api.app.core.clock import local_today, utc_now_naive
from api.app.schemas.checkins import CheckinSubmission, CheckinSubmissionResult
from api.app.services.checkin_catalog import get_today_checkin
from api.app.services.demo_users import WORKBOOK_PATH
from api.app.services.question_curator import get_personalization


workbook_lock = Lock()
ENERGY_VALUES = {"Very low": 1, "Low": 2, "Moderate": 3, "Good": 4, "Very good": 5}


class DuplicateCheckinError(ValueError):
    """Raised when a user tries to submit more than once on the same date."""


def normalized_date(value):
    return value.date() if hasattr(value, "date") else value


def has_submitted_today(username: str) -> bool:
    if not WORKBOOK_PATH.exists():
        return False
    today = local_today()
    with workbook_lock:
        workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
        if "Check-in Submissions" not in workbook.sheetnames:
            workbook.close()
            return False
        audit = workbook["Check-in Submissions"]
        found = any(
            str(row[1] or "").strip().lower() == username.lower()
            and normalized_date(row[2]) == today
            for row in audit.iter_rows(min_row=2, min_col=1, max_col=3, values_only=True)
        )
        workbook.close()
        return found


def expand_table(sheet, last_column: str) -> None:
    """Keep the visible Excel table/filter range aligned with appended records."""
    if sheet.tables:
        first_table = next(iter(sheet.tables.values()))
        first_table.ref = f"A1:{last_column}{sheet.max_row}"


def validate_submission(username: str, submission: CheckinSubmission) -> None:
    form = get_today_checkin(personalization=get_personalization(username))
    if submission.form_version != form.form_version:
        raise ValueError("The check-in form has changed. Refresh and try again.")
    for section in form.sections:
        for question in section.questions:
            value = submission.answers.get(question.id)
            if question.required and (value is None or value == "" or value == []):
                raise ValueError(f"Missing required answer: {question.id}")
            if value is None:
                continue
            if question.type in {"number", "linear_scale"}:
                try:
                    number = float(value)
                except (TypeError, ValueError) as exc:
                    raise ValueError(f"Invalid numeric answer: {question.id}") from exc
                if question.min_value is not None and number < question.min_value:
                    raise ValueError(f"Answer below minimum: {question.id}")
                if question.max_value is not None and number > question.max_value:
                    raise ValueError(f"Answer above maximum: {question.id}")


def save_checkin(username: str, submission: CheckinSubmission) -> CheckinSubmissionResult:
    validate_submission(username, submission)
    today = local_today()
    answers = submission.answers
    mood = float(answers["overall_mood"])
    stress = float(answers["stress_level"])
    sleep = float(answers.get("sleep_hours") or 0)
    energy = float(ENERGY_VALUES.get(str(answers["energy_level"]), 3))
    connectedness = float(answers.get("social_connection") or 3)
    support = str(answers.get("additional_support") or "No action")
    note = str(answers.get("additional_context") or "Daily check-in")[:200]

    with workbook_lock:
        workbook = load_workbook(WORKBOOK_PATH)
        audit = workbook["Check-in Submissions"] if "Check-in Submissions" in workbook.sheetnames else None
        if audit is not None:
            already_submitted = any(
                str(row[1] or "").strip().lower() == username.lower()
                and normalized_date(row[2]) == today
                for row in audit.iter_rows(min_row=2, min_col=1, max_col=3, values_only=True)
            )
            if already_submitted:
                workbook.close()
                raise DuplicateCheckinError("You have already completed today's check-in. Come back tomorrow.")
        history = workbook["15-Day History"]
        existing_row = None
        for row in range(2, history.max_row + 1):
            row_date = history.cell(row, 1).value
            row_username = str(history.cell(row, 2).value or "").lower()
            row_day = normalized_date(row_date)
            if row_username == username.lower() and row_day == today:
                existing_row = row
                break
        values = [today, username, mood, stress, sleep, energy, connectedness, True, support, note]
        if existing_row:
            for column, value in enumerate(values, 1):
                history.cell(existing_row, column, value)
        else:
            history.append(values)
        expand_table(history, "J")

        if audit is None:
            audit = workbook.create_sheet("Check-in Submissions")
            audit.append(["Submitted at (UTC)", "Username", "Date", "Form version", "Answers JSON"])
            audit.freeze_panes = "A2"
        audit.append([
            utc_now_naive(), username, today,
            submission.form_version, json.dumps(answers, ensure_ascii=False),
        ])
        expand_table(audit, "E")
        workbook.save(WORKBOOK_PATH)
        workbook.close()
    return CheckinSubmissionResult(
        date=today,
        username=username,
        created=existing_row is None,
        message="Your check-in was saved.",
    )
