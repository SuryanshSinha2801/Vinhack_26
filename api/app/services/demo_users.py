from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from api.app.core.security import hash_password
from api.app.database import SessionLocal
from api.app.models import DailyCheckin, User


WORKBOOK_PATH = Path(__file__).resolve().parents[3] / "data" / "demo" / "mindtrail_demo_6_months.xlsx"


def seed_demo_users() -> None:
    """Import every explicitly fictional demo account from the workbook."""
    if not WORKBOOK_PATH.exists():
        return
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    sheet = workbook["Demo Users"]
    with SessionLocal() as db:
        for username, password, display_name, _cohort, _use in sheet.iter_rows(
            min_row=6, min_col=1, max_col=5, values_only=True
        ):
            if not username or not password:
                continue
            username = str(username).strip().lower()
            if not db.scalar(select(User).where(User.username == username)):
                db.add(
                    User(
                        username=username,
                        display_name=str(display_name),
                        password_hash=hash_password(str(password)),
                    )
                )
        db.commit()
    workbook.close()


def seed_demo_history() -> None:
    """Import synthetic workbook history once into the relational database."""
    if not WORKBOOK_PATH.exists():
        return
    with SessionLocal() as db:
        if db.scalar(select(DailyCheckin.id).limit(1)) is not None:
            return
        users = {user.username: user.id for user in db.scalars(select(User)).all()}
        workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
        sheet = workbook["15-Day History"]
        rows: list[DailyCheckin] = []
        for values in sheet.iter_rows(min_row=2, min_col=1, max_col=10, values_only=True):
            record_date, username, mood, stress, sleep, energy, connectedness, completed, support, note = values
            user_id = users.get(str(username or "").strip().lower())
            if not user_id or not record_date:
                continue
            rows.append(
                DailyCheckin(
                    user_id=user_id,
                    checkin_date=record_date.date() if hasattr(record_date, "date") else record_date,
                    form_version="synthetic-seed-v1",
                    mood=float(mood or 0),
                    stress=float(stress or 0),
                    sleep_hours=float(sleep or 0),
                    energy=float(energy or 0),
                    connectedness=float(connectedness or 0),
                    checkin_completed=bool(completed),
                    support_response=str(support or "No action"),
                    note=str(note or "Synthetic demo record"),
                    answers={},
                )
            )
        db.add_all(rows)
        db.commit()
        workbook.close()
