from __future__ import annotations

import hashlib
import json
from datetime import date

from openpyxl import load_workbook

from api.app.core.clock import local_today
from api.app.schemas.checkins import CheckinPersonalization
from api.app.services.demo_users import WORKBOOK_PATH


STRATEGY_VERSION = "adaptive-checkin-v1"
FOCUS_LABELS = {
    "sleep": "sleep and recovery",
    "academic": "study pressure",
    "connection": "social connection",
    "support": "support preferences",
}


def _mean(values: list[float]) -> float:
    return sum(values) / len(values) if values else 0.0


def _slope(values: list[float]) -> float:
    if len(values) < 2:
        return 0.0
    center = (len(values) - 1) / 2
    denominator = sum((index - center) ** 2 for index in range(len(values)))
    return sum((index - center) * (value - _mean(values)) for index, value in enumerate(values)) / denominator


def rank_focus_areas(records: list[dict[str, float]], username: str, today: date) -> list[str]:
    """Rank non-clinical follow-up topics from recent wellbeing patterns."""
    if not records:
        rotation = ["sleep", "academic", "connection", "support"]
        digest = hashlib.sha256(f"{username}:{today.isoformat()}".encode()).digest()[0]
        start = digest % len(rotation)
        return [rotation[start], rotation[(start + 1) % len(rotation)]]

    recent = records[-14:]
    mood = [row["mood"] for row in recent]
    stress = [row["stress"] for row in recent]
    sleep = [row["sleep_hours"] for row in recent]
    energy = [row["energy"] for row in recent]
    connection = [row["connectedness"] for row in recent]
    scores = {
        "sleep": max(0.0, 7.0 - _mean(sleep)) + max(0.0, -_slope(sleep)) * 2,
        "academic": max(0.0, _mean(stress) - 2.5) + max(0.0, _slope(stress)) * 2,
        "connection": max(0.0, 3.5 - _mean(connection)) + max(0.0, -_slope(connection)),
        "support": max(0.0, 3.2 - _mean(mood)) + max(0.0, 3.2 - _mean(energy)),
    }
    tie_break = {name: index for index, name in enumerate(["sleep", "academic", "connection", "support"])}
    return sorted(scores, key=lambda name: (-scores[name], tie_break[name]))[:2]


def get_personalization(username: str) -> CheckinPersonalization:
    records: list[dict[str, float]] = []
    consent: str | None = None
    if WORKBOOK_PATH.exists():
        workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
        history = workbook["15-Day History"]
        for row in history.iter_rows(min_row=2, min_col=1, max_col=10, values_only=True):
            if str(row[1] or "").strip().lower() != username.lower():
                continue
            records.append({
                "date": row[0],
                "mood": float(row[2] or 0),
                "stress": float(row[3] or 0),
                "sleep_hours": float(row[4] or 0),
                "energy": float(row[5] or 0),
                "connectedness": float(row[6] or 0),
            })
        records.sort(key=lambda row: row["date"])
        if "Check-in Submissions" in workbook.sheetnames:
            audit = workbook["Check-in Submissions"]
            for row in audit.iter_rows(min_row=2, min_col=2, max_col=5, values_only=True):
                if str(row[0] or "").strip().lower() != username.lower() or not row[3]:
                    continue
                try:
                    answers = json.loads(str(row[3]))
                    consent = answers.get("trend_suggestions_consent", consent)
                except (TypeError, ValueError, json.JSONDecodeError):
                    continue
        workbook.close()

    allowed = consent != "No"
    focus_areas = rank_focus_areas(records if allowed else [], username, local_today())
    labels = [FOCUS_LABELS[area] for area in focus_areas]
    if records and allowed:
        explanation = f"Today includes {labels[0]} and {labels[1]} based on your recent check-in patterns."
    elif consent == "No":
        explanation = f"Today includes a rotating mix of {labels[0]} and {labels[1]}; recent trends were not used."
    else:
        explanation = f"Today starts with {labels[0]} and {labels[1]}; future forms can adapt as your history grows."
    return CheckinPersonalization(
        strategy_version=STRATEGY_VERSION,
        personalized=bool(records and allowed),
        focus_areas=focus_areas,
        based_on_days=min(len(records), 14) if allowed else 0,
        explanation=explanation,
    )
