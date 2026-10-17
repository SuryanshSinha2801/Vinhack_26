from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pathlib import Path
from shutil import copy2
from datetime import date
import json

from api.app.main import app
from api.app.services.checkin_store import normalized_date
from api.app.services.question_curator import get_personalization, rank_focus_areas


client = TestClient(app)


def login(test_client: TestClient) -> None:
    assert test_client.post(
        "/v1/auth/login", json={"username": "arjun01", "password": "Trail@Arjun26"}
    ).status_code == 200


def submission_payload() -> dict:
    return {
        "form_version": "2026-10-17",
        "answers": {
            "year_of_study": "Year 2",
            "overall_mood": 4,
            "stress_level": 2,
            "energy_level": "Good",
            "social_connection": 4,
            "sleep_quality": "Good",
            "sleep_hours": 7,
            "academic_workload": "Manageable",
            "additional_support": "No, I am doing okay",
            "immediate_help": "No",
            "follow_up_requested": "No",
            "trend_suggestions_consent": "Yes",
            "privacy_acknowledgement": [
                "I understand that this check-in is voluntary and is intended to support my wellbeing."
            ],
        },
    }


def test_today_checkin_matches_form_structure() -> None:
    assert TestClient(app).get("/v1/checkins/today").status_code == 401
    login(client)
    response = client.get("/v1/checkins/today")

    assert response.status_code == 200
    payload = response.json()
    assert payload["error"] is None

    checkin = payload["data"]
    assert checkin["title"] == "MindTrail – Student Wellbeing Check-in"
    assert isinstance(checkin["completed"], bool)
    assert len(checkin["sections"]) == 6
    assert checkin["personalization"]["strategy_version"] == "adaptive-checkin-v1"
    assert len(checkin["personalization"]["focus_areas"]) == 2

    questions = {
        question["id"]: question
        for section in checkin["sections"]
        for question in section["questions"]
    }
    assert 10 <= len(questions) <= 13
    assert questions["overall_mood"]["min_value"] == 1
    assert questions["overall_mood"]["max_value"] == 5
    assert questions["stress_level"]["max_label"] == "Very high"
    assert "student_id" not in questions
    assert questions["privacy_acknowledgement"]["required"] is True


def test_curator_prioritizes_sleep_and_academic_pressure() -> None:
    records = [
        {
            "mood": 3.0,
            "stress": 4.2 + index * 0.05,
            "sleep_hours": 5.5 - index * 0.05,
            "energy": 3.0,
            "connectedness": 4.0,
        }
        for index in range(14)
    ]
    focus = rank_focus_areas(records, "student", date(2026, 10, 17))
    assert set(focus) == {"sleep", "academic"}


def test_curator_respects_trend_personalization_opt_out(tmp_path, monkeypatch) -> None:
    source = Path("data/demo/mindtrail_demo_6_months.xlsx")
    target = tmp_path / "consent-demo.xlsx"
    copy2(source, target)
    workbook = load_workbook(target)
    workbook["Check-in Submissions"].append([
        None,
        "arjun01",
        date.today(),
        "2026-10-17",
        json.dumps({"trend_suggestions_consent": "No"}),
    ])
    workbook.save(target)
    workbook.close()
    monkeypatch.setattr("api.app.services.question_curator.WORKBOOK_PATH", target)

    personalization = get_personalization("arjun01")
    assert personalization.personalized is False
    assert personalization.based_on_days == 0
    assert "were not used" in personalization.explanation


def test_immediate_help_option_triggers_support_path() -> None:
    login(client)
    response = client.get("/v1/checkins/today")
    sections = response.json()["data"]["sections"]
    immediate_help = next(
        question
        for section in sections
        for question in section["questions"]
        if question["id"] == "immediate_help"
    )

    triggering_options = [
        option for option in immediate_help["options"] if option["triggers_immediate_support"]
    ]
    assert triggering_options == [
        {
            "value": "Yes – I need immediate help",
            "label": "Yes – I need immediate help",
            "triggers_immediate_support": True,
        }
    ]


def test_authenticated_checkin_is_saved_to_excel(tmp_path, monkeypatch) -> None:
    source = Path("data/demo/mindtrail_demo_6_months.xlsx")
    target = tmp_path / "demo.xlsx"
    copy2(source, target)
    monkeypatch.setattr("api.app.services.checkin_store.WORKBOOK_PATH", target)
    monkeypatch.setattr("api.app.services.question_curator.WORKBOOK_PATH", target)
    test_day = date(2099, 1, 2)
    monkeypatch.setattr("api.app.services.checkin_store.local_today", lambda: test_day)

    client = TestClient(app)
    login(client)
    assert client.get("/v1/checkins/today").json()["data"]["completed"] is False
    response = client.post("/v1/checkins", json=submission_payload())
    assert response.status_code == 200
    assert response.json()["data"]["username"] == "arjun01"
    assert response.json()["data"]["date"] == test_day.isoformat()
    assert client.get("/v1/checkins/today").json()["data"]["completed"] is True

    duplicate = client.post("/v1/checkins", json=submission_payload())
    assert duplicate.status_code == 409
    assert "already completed" in duplicate.json()["detail"]

    workbook = load_workbook(target, read_only=True, data_only=True)
    assert "Check-in Submissions" in workbook.sheetnames
    audit = workbook["Check-in Submissions"]
    assert audit.cell(audit.max_row, 2).value == "arjun01"
    submitted_rows = [
        row for row in audit.iter_rows(min_row=2, min_col=1, max_col=5, values_only=True)
        if row[1] == "arjun01" and normalized_date(row[2]) == test_day
    ]
    assert len(submitted_rows) == 1
    workbook.close()
